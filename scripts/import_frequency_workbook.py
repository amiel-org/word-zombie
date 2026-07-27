from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EXPECTED_SHEET_COUNTS = {
    "高频词分组": 660,
    "中频核心词分组": 660,
    "中频提升词分租": 660,
    "低频拓展词": 1_200,
}

SHEET_METADATA = {
    "高频词分组": ("high", None),
    "中频核心词分组": ("medium", "core"),
    "中频提升词分租": ("medium", "advanced"),
    "低频拓展词": ("low", None),
}

CHINESE_DIGITS = {
    "一": 1,
    "二": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "七": 7,
    "八": 8,
    "九": 9,
}


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_english(value: Any) -> str:
    return normalize_text(value).lower()


def parse_chinese_number(value: str) -> int:
    if value == "十":
        return 10
    if value.startswith("十"):
        return 10 + CHINESE_DIGITS[value[1:]]
    if value.endswith("十"):
        return CHINESE_DIGITS[value[0]] * 10
    if "十" in value:
        tens, ones = value.split("十", 1)
        return CHINESE_DIGITS[tens] * 10 + CHINESE_DIGITS[ones]
    return CHINESE_DIGITS[value]


def parse_group_number(header: Any) -> int | None:
    text = normalize_text(header)
    match = re.search(r"第([一二三四五六七八九十]+)组", text)
    return parse_chinese_number(match.group(1)) if match else None


def stable_word_id(english: str) -> str:
    digest = hashlib.sha1(normalize_english(english).encode("utf-8")).hexdigest()[:12]
    return f"freq-{digest}"


def extract_sheet_words(worksheet: Any) -> list[dict[str, Any]]:
    band, medium_band = SHEET_METADATA[worksheet.title]
    words: list[dict[str, Any]] = []

    for start_column in (1, 4, 7):
        headers: list[tuple[int, int]] = []
        for row_number in range(1, worksheet.max_row + 1):
            group_number = parse_group_number(worksheet.cell(row_number, start_column).value)
            if group_number is not None:
                headers.append((row_number, group_number))

        for header_index, (header_row, group_number) in enumerate(headers):
            end_row = (
                headers[header_index + 1][0]
                if header_index + 1 < len(headers)
                else worksheet.max_row + 1
            )
            group_words: list[dict[str, Any]] = []
            for row_number in range(header_row + 2, end_row):
                serial_value = worksheet.cell(row_number, start_column).value
                english = normalize_text(worksheet.cell(row_number, start_column + 1).value)
                meaning_zh = normalize_text(worksheet.cell(row_number, start_column + 2).value)
                if not isinstance(serial_value, (int, float)) or not english or not meaning_zh:
                    continue
                serial = int(serial_value)
                group_words.append(
                    {
                        "id": stable_word_id(english),
                        "english": english,
                        "meaningZh": meaning_zh,
                        "band": band,
                        **({"mediumBand": medium_band} if medium_band else {}),
                        "groupNumber": group_number,
                        "serial": serial,
                        "sourceSheet": worksheet.title,
                    }
                )

            if len(group_words) != 60:
                raise ValueError(
                    f"{worksheet.title} 第{group_number}组应有 60 词，实际 {len(group_words)}"
                )
            words.extend(group_words)

    words.sort(key=lambda word: (word["groupNumber"], word["serial"]))
    return words


def import_workbook(source: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    workbook = load_workbook(source, read_only=False, data_only=True)
    if workbook.sheetnames != list(EXPECTED_SHEET_COUNTS):
        raise ValueError(
            "Excel sheet 顺序或名称不符合预期：" + ", ".join(workbook.sheetnames)
        )

    by_sheet: dict[str, list[dict[str, Any]]] = {}
    for sheet_name in EXPECTED_SHEET_COUNTS:
        words = extract_sheet_words(workbook[sheet_name])
        expected = EXPECTED_SHEET_COUNTS[sheet_name]
        if len(words) != expected:
            raise ValueError(f"{sheet_name} 应有 {expected} 词，实际 {len(words)}")
        by_sheet[sheet_name] = words

    ordered_words = [
        *by_sheet["高频词分组"],
        *by_sheet["中频核心词分组"],
        *by_sheet["中频提升词分租"],
        *by_sheet["低频拓展词"],
    ]
    for source_order, word in enumerate(ordered_words, start=1):
        word["sourceOrder"] = source_order

    normalized_english = [normalize_english(word["english"]) for word in ordered_words]
    duplicate_english = sorted(
        word for word in set(normalized_english) if normalized_english.count(word) > 1
    )
    duplicate_ids = sorted(
        word_id
        for word_id in {word["id"] for word in ordered_words}
        if sum(item["id"] == word_id for item in ordered_words) > 1
    )
    malformed = [
        word["sourceOrder"]
        for word in ordered_words
        if not word["id"] or not word["english"] or not word["meaningZh"]
    ]
    if duplicate_english:
        raise ValueError(f"发现重复英文词条：{duplicate_english[:10]}")
    if duplicate_ids:
        raise ValueError(f"发现稳定 ID 冲突：{duplicate_ids[:10]}")
    if malformed:
        raise ValueError(f"发现必填字段缺失，sourceOrder={malformed[:10]}")
    if len(ordered_words) != 3_180:
        raise ValueError(f"总词数应为 3180，实际 {len(ordered_words)}")

    source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
    payload = {
        "schemaVersion": 1,
        "sourceFile": source.name,
        "sourceSha256": source_sha256,
        "total": len(ordered_words),
        "counts": {
            "high": 660,
            "medium": 1_320,
            "mediumCore": 660,
            "mediumAdvanced": 660,
            "low": 1_200,
        },
        "words": ordered_words,
    }
    audit = {
        "sourceFile": source.name,
        "sourceSha256": source_sha256,
        "sheetCounts": {name: len(words) for name, words in by_sheet.items()},
        "bandCounts": payload["counts"],
        "total": len(ordered_words),
        "uniqueNormalizedEnglish": len(set(normalized_english)),
        "duplicateEnglish": duplicate_english,
        "duplicateIds": duplicate_ids,
        "malformedSourceOrders": malformed,
    }
    return payload, audit


def main() -> None:
    parser = argparse.ArgumentParser(description="Import the senior frequency workbook")
    parser.add_argument("source", type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("data/senior_frequency_words.json"),
    )
    parser.add_argument(
        "--audit",
        type=Path,
        default=Path("data/senior_frequency_audit.json"),
    )
    args = parser.parse_args()

    payload, audit = import_workbook(args.source.resolve())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.audit.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    args.audit.write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(
        f"Imported {payload['total']} words: "
        f"high={payload['counts']['high']}, "
        f"medium={payload['counts']['medium']}, "
        f"low={payload['counts']['low']}"
    )


if __name__ == "__main__":
    main()
